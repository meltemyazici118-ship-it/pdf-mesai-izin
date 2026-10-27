
import io
import re
from datetime import datetime

import streamlit as st
import pandas as pd
import fitz
from PIL import Image, ImageOps, ImageEnhance
import pytesseract
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

st.set_page_config(
    page_title="PDF İzin–Mesai Kontrol",
    page_icon="📄",
    layout="wide"
)
st.title("📄 PDF İzin–Mesai Kontrol")
st.info(
    "Ücretsiz sürüm. PDF metni ve OCR ile bilgileri okumayı dener. "
    "El yazısı ve işaretli kutular mutlaka kontrol edilmelidir."
)
st.warning(
    "Bu sürüm PDKS 131 kayıtlarıyla karşılaştırma yapmaz. "
    "Personel belgeleri için kurumunuzun veri güvenliği kurallarına uyun."
)

COLS = [
    "Dosya", "Sayfa", "Ad Soyad", "Sicil No",
    "Belge Türü", "Talep", "Tarih", "Başlangıç",
    "Bitiş", "Açıklama", "Okuma Güveni"
]

DATE_RE = re.compile(
    r"(?<!\d)(\d{1,2})\s*[./-]\s*(\d{1,2})\s*[./-]\s*(20\d{2})(?!\d)"
)

def dates_in(text):
    found = []
    for d, m, y in DATE_RE.findall(text):
        try:
            date = datetime(int(y), int(m), int(d))
            value = date.strftime("%d.%m.%Y")
            if value not in found:
                found.append(value)
        except ValueError:
            pass
    return found

def ocr_page(page):
    pix = page.get_pixmap(
        matrix=fitz.Matrix(2.5, 2.5),
        alpha=False
    )
    img = Image.open(
        io.BytesIO(pix.tobytes("png"))
    ).convert("RGB")
    gray = ImageEnhance.Contrast(
        ImageOps.grayscale(img)
    ).enhance(1.8)

    attempts = []
    for picture, psm in [(img, 6), (gray, 11)]:
        try:
            result = pytesseract.image_to_string(
                picture,
                lang="tur+eng",
                config=f"--psm {psm}"
            )
        except pytesseract.TesseractError:
            result = pytesseract.image_to_string(
                picture,
                lang="eng",
                config=f"--psm {psm}"
            )
        attempts.append(result)

    return max(
        attempts,
        key=lambda text: (
            len(dates_in(text)),
            len(text.strip())
        )
    )

def parse_pdf(file):
    doc = fitz.open(
        stream=file.getvalue(),
        filetype="pdf"
    )
    if len(doc) > 20:
        raise ValueError("20 sayfadan uzun PDF'yi bölün.")

    rows = []
    for number, page in enumerate(doc, 1):
        direct = page.get_text("text").strip()
        method = "PDF metni" if len(direct) >= 40 else "OCR"
        text = direct if method == "PDF metni" else ocr_page(page)

        dates = dates_in(text)
        name = ""
        sicil = ""

        match = re.search(
            r"(?:Ad[ıi]\s*Soyad[ıi]|Ad[ıi]\s*ve\s*Soyad[ıi])"
            r"\s*[:：-]?\s*([^\n]{3,55})",
            text,
            re.I
        )
        if match:
            name = match.group(1).strip()

        match = re.search(
            r"Sicil\s*(?:No|Numaras[ıi])?"
            r"\s*[:：-]?\s*(\d{3,12})",
            text,
            re.I
        )
        if match:
            sicil = match.group(1)

        for date in dates or [""]:
            rows.append({
                "Dosya": file.name,
                "Sayfa": number,
                "Ad Soyad": name,
                "Sicil No": sicil,
                "Belge Türü": "",
                "Talep": "KONTROL EDİLEMEDİ",
                "Tarih": date,
                "Başlangıç": "",
                "Bitiş": "",
                "Açıklama": (
                    f"{method}; bilgileri PDF görüntüsünden doğrulayın."
                ),
                "Okuma Güveni": (
                    "DÜŞÜK" if method == "OCR" else "ORTA"
                )
            })
    return rows

files = st.file_uploader(
    "PDF belgelerini yükleyin",
    type=["pdf"],
    accept_multiple_files=True
)

if "rows" not in st.session_state:
    st.session_state.rows = []

if st.button(
    "🔍 PDF belgelerini oku",
    type="primary",
    disabled=not files
):
    all_rows = []
    for file in files:
        try:
            all_rows.extend(parse_pdf(file))
        except Exception as error:
            all_rows.append({
                "Dosya": file.name,
                "Sayfa": "",
                "Talep": "KONTROL EDİLEMEDİ",
                "Açıklama": f"Okuma hatası: {error}",
                "Okuma Güveni": "DÜŞÜK"
            })
    st.session_state.rows = all_rows

if files:
    with st.expander(
        "📄 PDF sayfasını gör ve bilgileri kontrol et",
        expanded=True
    ):
        selected = st.selectbox(
            "Belge",
            list(range(len(files))),
            format_func=lambda i: files[i].name
        )
        try:
            doc = fitz.open(
                stream=files[selected].getvalue(),
                filetype="pdf"
            )
            page_num = st.number_input(
                "Sayfa",
                min_value=1,
                max_value=len(doc),
                value=1,
                step=1
            )
            pix = doc[page_num - 1].get_pixmap(
                matrix=fitz.Matrix(1.5, 1.5),
                alpha=False
            )
            st.image(
                pix.tobytes("png"),
                caption=f"Sayfa {page_num}",
                use_container_width=True
            )
        except Exception as error:
            st.error(str(error))

if st.session_state.rows:
    st.subheader("Sonuçları PDF'ye bakarak düzeltin")
    st.caption(
        "Eksik tarihleri yeni satır ekleyerek girebilirsiniz. "
        "Saatler yanlış eşleştirmeyi önlemek için boş bırakılır."
    )

    frame = pd.DataFrame(
        st.session_state.rows
    ).reindex(columns=COLS).fillna("")

    edited = st.data_editor(
        frame,
        key="review_editor",
        num_rows="dynamic",
        hide_index=True,
        use_container_width=True,
        column_config={
            "Talep": st.column_config.SelectboxColumn(
                "Talep",
                options=[
                    "İZİN",
                    "MESAİ",
                    "KONTROL EDİLEMEDİ"
                ],
                required=True
            )
        }
    )

    def excel_bytes(data):
        output = io.BytesIO()
        with pd.ExcelWriter(
            output,
            engine="openpyxl"
        ) as writer:
            data.to_excel(
                writer,
                index=False,
                sheet_name="PDF Talepleri"
            )
            ws = writer.sheets["PDF Talepleri"]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions

            for cell in ws[1]:
                cell.fill = PatternFill(
                    "solid",
                    fgColor="17365D"
                )
                cell.font = Font(
                    bold=True,
                    color="FFFFFF"
                )

            colors = {
                "İZİN": "DFF0D8",
                "MESAİ": "DDEBF7",
                "KONTROL EDİLEMEDİ": "FFF2CC"
            }

            for row in ws.iter_rows(min_row=2):
                color = colors.get(
                    row[5].value,
                    "FFFFFF"
                )
                for cell in row:
                    cell.fill = PatternFill(
                        "solid",
                        fgColor=color
                    )
                    cell.alignment = Alignment(
                        vertical="top",
                        wrap_text=True
                    )

            widths = [
                35, 10, 25, 15, 25, 23,
                16, 15, 15, 68, 18
            ]
            for i, width in enumerate(widths, 1):
                ws.column_dimensions[
                    get_column_letter(i)
                ].width = width

        return output.getvalue()

    st.download_button(
        "⬇️ Kontrol edilen renkli Excel'i indir",
        data=excel_bytes(edited),
        file_name="PDF_Izin_Mesai_Kontrol.xlsx",
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        type="primary"
    )
